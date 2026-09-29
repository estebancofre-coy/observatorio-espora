"""Genera la presentación «Entornos Alimentarios en Coyhaique» (ESPORA).

Uso: python presentacion/generar_presentacion.py
Requiere: pip install python-pptx
"""
from pathlib import Path
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Entornos_Alimentarios_Coyhaique_ESPORA.pptx"
LOGO = ROOT.parent / "assets" / "logo-uaysen.png"

INK = RGBColor(0x18, 0x33, 0x46)
INK_SOFT = RGBColor(0x3F, 0x59, 0x66)
TEAL = RGBColor(0x0F, 0x71, 0x80)
DEEP = RGBColor(0x0F, 0x5D, 0x74)
PAPER = RGBColor(0xF3, 0xF7, 0xF4)
MINT = RGBColor(0xE5, 0xF2, 0xEE)
LINE = RGBColor(0xD5, 0xE1, 0xDF)
PALE = RGBColor(0x9F, 0xD3, 0xD0)
MIST = RGBColor(0xD8, 0xEA, 0xE8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
W, H = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]
TOTAL = 11


def rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def add_runs(paragraph, content, size, color, bold=False):
    """Escribe texto; *así* se convierte en cursiva (necesario para APA 7)."""
    for i, part in enumerate(re.split(r"\*", content)):
        if not part:
            continue
        r = paragraph.add_run()
        r.text = part
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.bold = bold
        r.font.italic = i % 2 == 1


def text(slide, x, y, w, h, lines, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=6, line_spacing=1.15):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        p.line_spacing = line_spacing
        add_runs(p, line, size, color, bold)
    return tb


def base(title, kicker, number):
    slide = prs.slides.add_slide(BLANK)
    rect(slide, 0, 0, W, H, PAPER)
    rect(slide, 0, 0, Inches(0.22), H, TEAL)
    text(slide, Inches(0.8), Inches(0.55), Inches(10), Inches(0.35), kicker.upper(),
         size=12, color=TEAL, bold=True)
    text(slide, Inches(0.8), Inches(0.9), Inches(10.3), Inches(0.9), title, size=32, bold=True)
    text(slide, Inches(0.8), Inches(6.95), Inches(9), Inches(0.3),
         "Entornos Alimentarios en Coyhaique · Proyecto ESPORA", size=10, color=INK_SOFT)
    text(slide, Inches(11.5), Inches(6.95), Inches(1.1), Inches(0.3), f"{number} / {TOTAL}",
         size=10, color=INK_SOFT, align=PP_ALIGN.RIGHT)
    return slide


def card(slide, x, y, w, h, heading, body, size=15, fill=WHITE):
    rect(slide, x, y, w, h, fill)
    rect(slide, x, y, w, Inches(0.06), TEAL)
    text(slide, x + Inches(0.3), y + Inches(0.3), w - Inches(0.6), Inches(0.5), heading,
         size=17, bold=True, color=DEEP)
    text(slide, x + Inches(0.3), y + Inches(0.85), w - Inches(0.6), h - Inches(1.05), body,
         size=size, color=INK_SOFT)


# 1. Portada
s = prs.slides.add_slide(BLANK)
rect(s, 0, 0, W, H, INK)
rect(s, Inches(8.9), 0, Inches(4.433), H, DEEP)
for cx, cy, d in [(10.2, 1.3, 1.6), (11.6, 2.9, 0.9), (9.7, 3.6, 0.55), (11.2, 4.6, 1.25), (10.0, 5.7, 0.7)]:
    rect(s, Inches(cx), Inches(cy), Inches(d), Inches(d), TEAL, MSO_SHAPE.OVAL)
text(s, Inches(0.8), Inches(0.8), Inches(7.6), Inches(0.4), "PROYECTO ESPORA · COYHAIQUE, AYSÉN",
     size=13, color=PALE, bold=True)
text(s, Inches(0.8), Inches(1.5), Inches(7.8), Inches(2.4), "Entornos Alimentarios en Coyhaique",
     size=50, color=WHITE, bold=True, line_spacing=1.0)
text(s, Inches(0.8), Inches(3.7), Inches(7.6), Inches(1.0),
     "Observación territorial de almacenes de barrio y minimarkets: disponibilidad, precios y origen de los alimentos",
     size=19, color=MIST)
text(s, Inches(0.8), Inches(5.05), Inches(7.8), Inches(1.4), [
    "Investigación original: PhD. Ana Zazo Moratalla, Universidad Rey Juan Carlos",
    "Aplicación: Dr. Esteban Cofré-Morales, Universidad de Aysén",
    "Marco de referencia alimentaria: FAO",
], size=14, color=WHITE, spacing=4)
text(s, Inches(0.8), Inches(6.7), Inches(7.8), Inches(0.35),
     "ESPORA: Economía Social, Precios, Orgánicos, Residuos y Alimentación · 2026",
     size=11, color=PALE)

# 2. Propósito
s = base("¿Qué alimentos están al alcance del barrio?", "Propósito", 2)
text(s, Inches(0.8), Inches(2.0), Inches(6.2), Inches(4.5), [
    "El comercio de proximidad define buena parte del entorno alimentario cotidiano. "
    "Este estudio observa, local por local, qué se ofrece, a qué precio y de dónde viene.",
    "El objetivo es describir los entornos alimentarios de Coyhaique a escala de unidad vecinal "
    "y aportar evidencia para la planificación urbana y la Economía Social y Solidaria (ESS).",
], size=19, color=INK, spacing=16)
for i, (h_, b_) in enumerate([
    ("Disponibilidad", "¿Qué alimentos y preparaciones saludables y no saludables se ofrecen?"),
    ("Precio", "¿Cuál es el rango de precios, del más bajo al más alto, de una canasta definida?"),
    ("Origen", "¿Qué proporción de proteínas, lácteos y huevos proviene del territorio?"),
]):
    card(s, Inches(7.6), Inches(1.95 + i * 1.6), Inches(4.95), Inches(1.4), h_, b_, size=14)

# 3. Marco
s = base("Marco de referencia", "Fundamentos", 3)
card(s, Inches(0.8), Inches(2.0), Inches(3.8), Inches(4.5), "Entornos alimentarios",
     ["La FAO entiende el entorno alimentario como el espacio donde las personas acceden a los "
      "alimentos y deciden qué comer; influir en él es una vía para dietas saludables.",
      "(Food and Agriculture Organization of the United Nations [FAO], 2016)"], size=14)
card(s, Inches(4.77), Inches(2.0), Inches(3.8), Inches(4.5), "Resiliencia alimentaria",
     ["La planificación urbana puede evaluar la seguridad y la resiliencia alimentaria de las "
      "ciudades chilenas con indicadores territoriales.",
      "(Zazo-Moratalla & Orellana-McBride, 2025)"], size=14)
card(s, Inches(8.74), Inches(2.0), Inches(3.8), Inches(4.5), "Pautas de observación",
     ["Las pautas de ambientes alimentarios desarrolladas en la Universidad del Bío-Bío orientan "
      "el registro de disponibilidad en puntos de venta.",
      "(Araneda-Flores et al., 2024)"], size=14)

# 4. Espora
s = base("La espora como metáfora de trabajo", "Identidad del proyecto", 4)
for i, (h_, b_) in enumerate([
    ("Dispersión", "El levantamiento se distribuye por las 34 unidades vecinales, con equipos que trabajan en paralelo."),
    ("Resistencia", "Los datos sobreviven sin conexión: cada ficha se guarda en el dispositivo hasta sincronizar."),
    ("Regeneración", "El conocimiento vuelve al territorio: compostaje, circuitos cortos y ESS como horizonte de uso."),
]):
    card(s, Inches(0.8 + i * 3.97), Inches(2.1), Inches(3.75), Inches(3.3), h_, b_, size=16)
text(s, Inches(0.8), Inches(5.8), Inches(11.7), Inches(0.8),
     "ESPORA reúne Economía Social, Precios, Orgánicos, Residuos y Alimentación en una sola mirada territorial.",
     size=16, color=INK_SOFT)

# 5. Territorio y muestra
s = base("Territorio y muestra", "Diseño", 5)
for i, (num, lab) in enumerate([("107", "locales en la muestra sugerida"),
                                 ("34", "unidades vecinales de la comuna"),
                                 ("216", "puntos en la capa KML de locales")]):
    x = Inches(0.8 + i * 3.97)
    text(s, x, Inches(2.0), Inches(3.7), Inches(1.2), num, size=64, color=TEAL, bold=True)
    text(s, x, Inches(3.25), Inches(3.7), Inches(0.6), lab, size=16, color=INK)
text(s, Inches(0.8), Inches(4.3), Inches(11.7), Inches(2.4), [
    "Cada local de la muestra se asigna a una unidad vecinal a partir de sus coordenadas, con los polígonos "
    "de unidades vecinales del Censo 2017 (Ministerio de Desarrollo Social y Familia, s.f.).",
    "La muestra y la capa de locales por subcategoría son materiales del proyecto (Proyecto ESPORA, 2026a, 2026c). "
    "Los límites censales pueden diferir de la división municipal vigente: los puntos cercanos a un borde se validan en terreno.",
], size=16, color=INK_SOFT, spacing=12)

# 6. Instrumentos
s = base("Cinco fichas, una visita", "Instrumentos en terreno", 6)
steps = [("Identificación", "Local, UV, coordenadas, persona recolectora y foto del frontis."),
         ("Clasificación", "Atención y surtido: almacén de barrio o minimarket."),
         ("Disponibilidad", "Pauta de alimentos y preparaciones saludables y no saludables."),
         ("Precios", "Canasta de 28 productos: precio más bajo y más alto (Proyecto ESPORA, 2026b)."),
         ("Origen", "Carnes, pescado, lácteos y huevos: ¿local o externo?")]
for i, (h_, b_) in enumerate(steps):
    x = Inches(0.8 + i * 2.39)
    rect(s, x, Inches(2.1), Inches(0.7), Inches(0.7), TEAL, MSO_SHAPE.OVAL)
    text(s, x, Inches(2.1), Inches(0.7), Inches(0.7), str(i + 1), size=20, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    if i < len(steps) - 1:
        rect(s, x + Inches(0.8), Inches(2.43), Inches(1.45), Inches(0.04), LINE)
    text(s, x, Inches(3.05), Inches(2.2), Inches(0.5), h_, size=17, bold=True, color=DEEP)
    text(s, x, Inches(3.6), Inches(2.15), Inches(2.4), b_, size=14, color=INK_SOFT)
text(s, Inches(0.8), Inches(6.1), Inches(11.7), Inches(0.6),
     "Todas las fichas comparten un ID de visita; cada instrumento se guarda por separado y puede retomarse.",
     size=14, color=INK)

# 7. Regla de clasificación
s = base("¿Almacén de barrio o minimarket?", "Regla de clasificación", 7)
rect(s, Inches(0.8), Inches(2.0), Inches(3.6), Inches(4.4), INK)
text(s, Inches(1.1), Inches(2.3), Inches(3.0), Inches(0.4), "CONDICIÓN OBLIGATORIA", size=12,
     color=PALE, bold=True)
text(s, Inches(1.1), Inches(2.8), Inches(3.0), Inches(1.0), "Autoservicio", size=30, color=WHITE, bold=True)
text(s, Inches(1.1), Inches(3.8), Inches(3.0), Inches(2.4),
     "El cliente toma los productos por sí mismo. Si todo se pide tras el mesón, es almacén de barrio.",
     size=15, color=MIST)
text(s, Inches(4.8), Inches(2.0), Inches(7.7), Inches(0.5), "…y al menos 2 de estos 3 criterios:",
     size=18, bold=True, color=INK)
for i, (h_, b_) in enumerate([
    ("Frutas y verduras", "Más de 10 variedades distintas (11 o más)."),
    ("Proteínas", "Al menos 2 categorías: vacuno, cerdo, pollo, cordero, pescados o mariscos."),
    ("Lácteos y huevos", "Oferta habitual de ambos."),
]):
    y = Inches(2.7 + i * 1.2)
    rect(s, Inches(4.8), y, Inches(7.7), Inches(1.0), WHITE)
    text(s, Inches(5.1), y + Inches(0.15), Inches(2.6), Inches(0.7), h_, size=16, bold=True, color=DEEP,
         anchor=MSO_ANCHOR.MIDDLE)
    text(s, Inches(7.8), y + Inches(0.15), Inches(4.5), Inches(0.7), b_, size=14, color=INK_SOFT,
         anchor=MSO_ANCHOR.MIDDLE)
text(s, Inches(0.8), Inches(6.5), Inches(11.7), Inches(0.4),
     "La preclasificación de la muestra no se muestra en terreno para no sesgar la observación.",
     size=13, color=INK_SOFT)

# 8. Flujo de datos
s = base("Del local a la base de datos", "Flujo de información", 8)
nodes = [("Dispositivo", "Formulario web en el teléfono"), ("Cola local", "Respaldo sin conexión"),
         ("Apps Script", "Valida y clasifica"),
         ("Google Sheets", "Visitas, Clasificación, Disponibilidad, Precios, Origen"),
         ("Mapa", "Consolidación y KML por UV (en diseño)")]
for i, (h_, b_) in enumerate(nodes):
    x = Inches(0.8 + i * 2.43)
    rect(s, x, Inches(2.3), Inches(2.1), Inches(2.4), MINT if i == 4 else WHITE)
    rect(s, x, Inches(2.3), Inches(2.1), Inches(0.06), TEAL)
    text(s, x + Inches(0.2), Inches(2.6), Inches(1.7), Inches(0.5), h_, size=16, bold=True, color=DEEP)
    text(s, x + Inches(0.2), Inches(3.15), Inches(1.75), Inches(1.5), b_, size=13, color=INK_SOFT)
    if i < len(nodes) - 1:
        rect(s, x + Inches(2.13), Inches(3.35), Inches(0.28), Inches(0.3), TEAL, MSO_SHAPE.RIGHT_ARROW)
text(s, Inches(0.8), Inches(5.1), Inches(11.7), Inches(1.5), [
    "El backend vuelve a validar cada envío y calcula la clasificación con la misma regla que la aplicación.",
    "Los códigos de locales nuevos incluyen las iniciales de quien registra (por ejemplo, LM108NMB) para evitar duplicados entre equipos.",
], size=15, color=INK_SOFT, spacing=8)

# 9. Calidad de datos
s = base("Calidad de datos desde el diseño", "Control", 9)
for i, (h_, b_) in enumerate([
    ("Observación sin sesgo", "El tipo registrado en la muestra no aparece en pantalla; la categoría la decide lo observado."),
    ("Autocapacitación", "Guía con simulador y casos de práctica antes de salir a terreno."),
    ("Estructura estable", "Las columnas nuevas se anexan al final; nunca se borran datos ni encabezados."),
    ("Trazabilidad", "ID de visita común, fotos en Drive y justificación obligatoria ante correcciones manuales."),
]):
    col, row = i % 2, i // 2
    card(s, Inches(0.8 + col * 6.0), Inches(2.0 + row * 2.3), Inches(5.75), Inches(2.05), h_, b_, size=15)

# 10. Próximos pasos
s = base("Estado y próximos pasos", "Hoja de ruta", 10)
text(s, Inches(0.8), Inches(2.0), Inches(5.6), Inches(0.5), "En curso", size=20, bold=True, color=DEEP)
text(s, Inches(0.8), Inches(2.6), Inches(5.6), Inches(3.8), [
    "Levantamiento en terreno por unidad vecinal.",
    "Registro de locales nuevos no incluidos en la muestra.",
    "Capacitación de estudiantes y profesionales de apoyo.",
], size=16, color=INK_SOFT, spacing=12)
text(s, Inches(7.0), Inches(2.0), Inches(5.6), Inches(0.5), "Siguiente", size=20, bold=True, color=DEEP)
text(s, Inches(7.0), Inches(2.6), Inches(5.6), Inches(3.8), [
    "Hoja maestra de locales vigentes con su clasificación observada.",
    "Exportación KML/GeoJSON para mapas por UV y por tipo de local.",
    "Análisis de disponibilidad, precios y origen por territorio.",
    "Validación metodológica de las definiciones de conteo.",
], size=16, color=INK_SOFT, spacing=12)

# 11. Referencias (APA 7, sangría francesa)
s = base("Referencias", "APA 7.ª edición", 11)
refs = [
    "Araneda-Flores, J., Toledo, Á., Inzunza, C., Córdova, C., & Pinheiro, A. C. (2024). Ambiente alimentario "
    "alrededor de establecimientos educacionales municipalizados de la ciudad de Chillán. *Revista Chilena de Nutrición, 51*(6), "
    "446–452. https://doi.org/10.4067/s0717-75182024000600446",
    "Cofré-Morales, E. (2026). *ESPORA Coyhaique: Plataforma de levantamiento de entornos alimentarios* [Aplicación web]. "
    "Universidad de Aysén. https://estebancofre-coy.github.io/observatorio-espora/",
    "Food and Agriculture Organization of the United Nations. (2016). *Influencing food environments for healthy diets*. "
    "https://openknowledge.fao.org/handle/20.500.14283/i6484en",
    "Ministerio de Desarrollo Social y Familia. (s.f.). *Unidades vecinales* [Conjunto de datos geoespaciales]. "
    "Infraestructura de Datos Espaciales de Chile. https://www.geoportal.cl/geoportal/catalog/36395/Unidades%20vecinales",
    "Proyecto ESPORA. (2026a). *Locales manuales por subcategoría* [Conjunto de datos no publicado, KML]. Universidad de Aysén.",
    "Proyecto ESPORA. (2026b). *Matriz de recolección de precios de Coyhaique* [Conjunto de datos no publicado]. Universidad de Aysén.",
    "Proyecto ESPORA. (2026c). *Muestra de locales de Coyhaique* [Conjunto de datos no publicado]. Universidad de Aysén.",
    "Zazo-Moratalla, A., & Orellana-McBride, A. (2025). How to assess urban food resilience? Moving towards food security "
    "in Chilean cities. *Sustainability, 17*(17), Artículo 7924. https://doi.org/10.3390/su17177924",
]
tb = text(s, Inches(0.8), Inches(1.85), Inches(11.8), Inches(5.0), refs, size=13, color=INK, spacing=6,
          line_spacing=1.05)
for p in tb.text_frame.paragraphs:
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(Inches(0.5))))
    pPr.set("indent", str(-int(Inches(0.5))))

if LOGO.exists():
    for slide in list(prs.slides)[1:]:
        slide.shapes.add_picture(str(LOGO), Inches(11.35), Inches(0.45), height=Inches(0.55))

props = prs.core_properties
props.title = "Entornos Alimentarios en Coyhaique"
props.author = "Ana Zazo Moratalla; Esteban Cofré-Morales"
props.subject = "Proyecto ESPORA: entornos alimentarios, precios y origen de alimentos en Coyhaique"
props.keywords = "entornos alimentarios; Coyhaique; Aysén; ESPORA; Economía Social y Solidaria"

prs.save(OUT)
print(f"Presentación generada: {OUT}")